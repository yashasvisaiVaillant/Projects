import math
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


class ExcelReportGenerator:
    def __init__(self, output_file):
        self.output_file = output_file

    def generate(self, test_name, failures, report_file):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Failures"

        worksheet.append(["Test Name", "Failure Step", "Failure Reason"])
        for cell in worksheet[1]:
            cell.font = Font(bold=True)

        report_uri = Path(report_file).resolve().as_uri()
        for index, (step, reason) in enumerate(failures):
            worksheet.append([
                test_name if index == 0 else '',
                step,
                reason
            ])
            reason_cell = worksheet.cell(row=index + 2, column=3)
            expected_text, separator, actual_text = reason.partition('\nActual: ')
            if separator:
                reason_cell.value = CellRichText(
                    f'{expected_text}\n',
                    TextBlock(
                        InlineFont(color='FFFF0000'),
                        f'Actual: {actual_text}'
                    )
                )
            else:
                reason_cell.value = CellRichText(
                    TextBlock(InlineFont(color='FFFF0000'), reason)
                )

            reason_cell.hyperlink = report_uri
            reason_cell.style = 'Hyperlink'

        ExcelFormatter.format_worksheet(worksheet)
        workbook.save(self.output_file)


class ExcelFormatter:
    @classmethod
    def format_combined_worksheet(cls, worksheet):
        cls.format_worksheet(worksheet)
        cls._merge_test_name_rows(worksheet)

    @classmethod
    def _merge_test_name_rows(cls, worksheet):
        group_start = None

        for row_index in range(2, worksheet.max_row + 1):
            test_name = worksheet.cell(row=row_index, column=1).value
            if test_name not in (None, ''):
                if group_start is not None:
                    cls._merge_test_name_group(
                        worksheet, group_start, row_index - 1
                    )
                group_start = row_index

        if group_start is not None:
            cls._merge_test_name_group(
                worksheet, group_start, worksheet.max_row
            )

    @staticmethod
    def _merge_test_name_group(worksheet, start_row, end_row):
        if end_row > start_row:
            worksheet.merge_cells(
                start_row=start_row,
                start_column=1,
                end_row=end_row,
                end_column=1
            )

        worksheet.cell(row=start_row, column=1).alignment = Alignment(
            vertical='center',
            wrap_text=True
        )

    @staticmethod
    def format_worksheet(worksheet):
        # Fixed caps keep long paths and diagnostics compact while wrapping.
        maximum_widths = {1: 35, 2: 70, 3: 65}

        for column_index, column_cells in enumerate(
            worksheet.columns, start=1
        ):
            max_length = 0
            column_letter = get_column_letter(column_index)
            for cell in column_cells:
                cell.alignment = Alignment(wrap_text=True, vertical='top')
                if cell.value:
                    line_length = max(
                        len(line) for line in str(cell.value).splitlines()
                    )
                    max_length = max(max_length, line_length)

            maximum_width = maximum_widths.get(column_index, 50)
            worksheet.column_dimensions[column_letter].width = min(
                max(max_length + 2, 12),
                maximum_width
            )

        worksheet.row_dimensions[1].height = 22
        # Estimate enough height for explicit and width-induced line wrapping.
        for row_index in range(2, worksheet.max_row + 1):
            wrapped_lines = 1
            for column_index in range(1, worksheet.max_column + 1):
                value = worksheet.cell(row=row_index, column=column_index).value
                if value is None:
                    continue

                width = worksheet.column_dimensions[
                    get_column_letter(column_index)
                ].width
                line_count = sum(
                    max(1, math.ceil(len(line) / width))
                    for line in str(value).splitlines()
                )
                wrapped_lines = max(wrapped_lines, line_count)

            worksheet.row_dimensions[row_index].height = min(
                wrapped_lines * 15,
                150
            )
