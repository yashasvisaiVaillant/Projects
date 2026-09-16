# This script generates a failure report from an HTML file containing test results.

import os
import pandas as pd
from bs4 import BeautifulSoup
from failure_traverser import FailureTraverser
from ExcelFormatter import ExcelFormatter


def read_html_with_detected_encoding(path): 
    # Encoding robustness: Why it’s needed:
    # It was observed that some of the Test data HTML files came in different encodings (UTF‑8, UTF‑16)
    # This can raise UnicodeDecodeError or silently corrupt characters.
    # A robust approach tries detection first, then falls back safely.
    
    # Read raw bytes 
    with open(path, 'rb') as f: 
        raw = f.read()
    # Try chardet if available
    enc = None
    try:
        import chardet
        info = chardet.detect(raw) or {}
        enc = info.get('encoding')
    except Exception:
        enc = None

    # Try detected encoding strictly; on error, fall back
    if enc:
        try:
            return raw.decode(enc, errors='strict')
        except Exception:
            pass

    # Try BOM-aware UTF-8 first
    try:
        return raw.decode('utf-8-sig', errors='strict')
    except Exception:
        # Final fallback that never fails (may replace undecodable bytes)
        return raw.decode('utf-8', errors='replace')

    
class FailureReportGenerator:
    def __init__(self, html_file, output_file):
        self.html_file = html_file
        self.output_file = output_file

    def run(self):
        # Extract test name from parent folder of the HTML file
        test_name = os.path.basename(os.path.dirname(os.path.abspath(self.html_file)))

        # Read HTML content (robust decoding)
        html_content = read_html_with_detected_encoding(self.html_file)

        # Parse HTML
        soup = BeautifulSoup(html_content, 'html.parser')

        # Find all 'div' elements with class 'Failed' anywhere in the document
        failed_divs = soup.find_all('div', class_='Failed')

        traverser = FailureTraverser()
        for div in failed_divs:
            traverser.traverse_failed_div(div)

        print(f"Total unique failures recorded: {len(traverser.failure_set)}")

        # Do not generate a report if there are no failures
        if not traverser.failure_set:
            print(f"No failures found in {self.html_file}. Skipping report generation.")
            return False
            
        # Build sorted list of failure steps
        steps = sorted(
            traverser.failure_set,
            key=lambda s: s.casefold() if isinstance(s, str) else str(s).casefold()
        )

        # Prepare data for DataFrame: first row with test name + first failure step
        rows = []

        if steps:
            # First row: include test name and first failure step
            rows.append([test_name, steps[0]])
            # Remaining rows: only failure steps, leave test name cell empty
            for step in steps[1:]:
                rows.append(['', step])
        else:
            # No failures, no report
            print(f"No failures found in {self.html_file}. Skipping report generation.")
            return False

        # Create DataFrame
        df = pd.DataFrame(rows, columns=['Test Name', 'Failure Step'])

        # Write to Excel with openpyxl
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = "Failures"

        # Write header
        headers = ["Test Name", "Failure Step"]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)

        # Write data rows
        for row in rows:
            ws.append(row)

        # Auto-size columns
        col_widths = {}
        for row in ws.iter_rows(values_only=True):
            for idx, value in enumerate(row, start=1):
                length = len(str(value)) if value is not None else 0
                if length > col_widths.get(idx, 0):
                    col_widths[idx] = length
        for idx, width in col_widths.items():
            ws.column_dimensions[get_column_letter(idx)].width = min(width + 2, 80)

        # Save workbook
        wb.save(self.output_file)
        print(f"Excel report written to: {self.output_file}")
        return True